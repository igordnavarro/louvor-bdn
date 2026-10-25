#!/usr/bin/env python3
"""Gera a planilha de controle do repositório (controle/) a partir de cifras/*.txt e dos pads.

Uso:  python3 tools/planilha.py
Saída: controle/Worship Church - Controle.xlsx  (abas "Músicas" e "Pads")
       controle/musicas.csv e controle/pads.csv  (mesmo conteúdo, fácil de ver no GitHub)

Sempre que adicionar/editar uma música: rode tools/build.py e depois este script.
"""
import csv, glob, json, os, re, sys, unicodedata, difflib
sys.path.insert(0, os.path.dirname(__file__))
from nomes import titulo, artista, nome_padrao

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'controle')
SITE = 'https://igordnavarro.github.io/louvor-bdn/'
REPO = 'https://github.com/igordnavarro/louvor-bdn/blob/main/'
NOTE = ['C','C#','D','D#','E','F','F#','G','G#','A','A#','B']
FLAT = {'Db':'C#','Eb':'D#','Gb':'F#','Ab':'G#','Bb':'A#'}

def read(f):
    meta, body = {}, []
    for line in open(f, encoding='utf-8').read().split('\n'):
        m = re.match(r'^@(\w+)\s?(.*)$', line)
        if m and not body: meta[m.group(1)] = m.group(2).strip(); continue
        body.append(line)
    meta['_body'] = '\n'.join(body).strip(); return meta

def norm(s):
    s = unicodedata.normalize('NFKD', (s or '').lower()); s = ''.join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r'\(.*?\)', '', s); s = re.sub(r'[^a-z0-9]+', ' ', s); return ' '.join(s.split())

def songs():
    ids = json.load(open(os.path.join(ROOT, 'tools', 'ids.json'), encoding='utf-8'))
    out = []
    for f in sorted(glob.glob(os.path.join(ROOT, 'cifras', '*.txt'))):
        fn = os.path.basename(f); m = read(f)
        t, a = titulo(m.get('title', '')), artista(m.get('artist', ''))
        bpm = re.match(r'\d+', m.get('bpm', ''))
        out.append(dict(id=ids.get(fn, ''), nome=nome_padrao(t, a), titulo=t, artista=a,
            tom=m.get('tom', ''), pad=m.get('pad', ''), bpm=int(bpm.group(0)) if bpm else '',
            capo=int(m['capo']) if m.get('capo', '').isdigit() and int(m['capo']) else '',
            cifra='Sim' if m['_body'] else 'Não', arquivo=m.get('arquivo', ''), slide=m.get('slide', ''),
            repo='cifras/' + fn, obs=''))
    marca_duplicadas(out)
    return out

def marca_duplicadas(L):
    """Preenche a coluna de observação com repetições e possíveis repetições."""
    def base(s): return norm(re.sub(r'(?i)simplificada', '', s))
    def add(x, txt):
        if txt not in x['obs']: x['obs'] = (x['obs'] + ' | ' if x['obs'] else '') + txt
    groups = {}
    for x in L: groups.setdefault(base(x['titulo']), []).append(x)
    for k, g in groups.items():
        if len(g) < 2: continue
        for x in g:
            for y in g:
                if x is y: continue
                ref = f"{y['titulo']} ({y['artista'] or 's/ artista'}, tom {y['tom'] or '?'}, id {y['id']})"
                if norm(x['artista']) == norm(y['artista']) or not x['artista'] or not y['artista']:
                    if 'simplificada' in (x['titulo'] + y['titulo']).lower() and ('simplificada' in x['titulo'].lower()) != ('simplificada' in y['titulo'].lower()):
                        add(x, 'Versão simplificada/completa de ' + ref)
                    elif x['tom'] == y['tom']:
                        add(x, 'REPETIDA: mesma música e mesmo tom que ' + ref)
                    else:
                        add(x, 'Mesma música em outro tom: ' + ref)
                else:
                    add(x, 'Mesmo nome, outro artista (provavelmente outra música): ' + ref)
    # nomes parecidos (erro de digitação/variação) e músicas que aparecem dentro de medleys
    keys = list(groups)
    for i, a in enumerate(keys):
        for b in keys[i + 1:]:
            if not a or not b: continue
            r = difflib.SequenceMatcher(None, a, b).ratio()
            if r >= 0.86:
                for x in groups[a]:
                    for y in groups[b]: add(x, f"Possível repetida (nome parecido): {y['titulo']} – {y['artista']} (id {y['id']})")
                for y in groups[b]:
                    for x in groups[a]: add(y, f"Possível repetida (nome parecido): {x['titulo']} – {x['artista']} (id {x['id']})")
    for x in L:
        if ',' in str(x['slide']):
            add(x, f"REPETIDA no PowerPoint: aparece nos slides {x['slide']} (no app ficou uma só)")
    for x in L:
        partes = [norm(p) for p in re.split(r'[/+]', x['titulo'])]
        if len(partes) < 2: continue
        for y in L:
            if y is x: continue
            if base(y['titulo']) in partes:
                add(y, f"Também aparece no medley: {x['titulo']} (id {x['id']})")
                add(x, f"Contém a música: {y['titulo']} – {y['artista']} (id {y['id']})")

def pads():
    rows = []
    builtin = [('A','A','Pad A'),('B','B','Pad B'),('C#','Cs','Pad C# (Db)'),('D','D','Pad D'),('E','E','Pad E'),('G','G','Pad G'),
               ('Cm','Cm','Pad Cm'),('Em','Em-2','Pad Em'),('Em','Em','Pad Em (variação)'),('Bm','Bm','Pad Bm')]
    for tom, f, nome in builtin:
        rows.append(dict(tom=tom, nome=nome, link=SITE + f'pads/{f}.m4a', fonte='Incluído no app (arquivo de áudio)',
                         offline='Sim', arquivo=f'pads/{f}.m4a'))
    src = open(os.path.join(ROOT, 'app.js'), encoding='utf-8').read()
    m = re.search(r'const YT_PADS = \{(.*?)\n\};', src, re.S)
    for name, ids in re.findall(r"name:'([^']+)', ids:\{([^}]*)\}", m.group(1)):
        for r, vid in re.findall(r"(\d+):'([\w-]+)'", ids):
            rows.append(dict(tom=NOTE[int(r)], nome=f'{name} em {NOTE[int(r)]}', link=f'https://www.youtube.com/watch?v={vid}',
                             fonte='YouTube', offline='Não', arquivo=''))
    rows.append(dict(tom='Vários', nome='Playlist de pads do ministério', fonte='YouTube (playlist)', offline='Não', arquivo='',
                     link='https://music.youtube.com/playlist?list=PLiwIL30R32_MHPvGvPSLmGi4Gs7k6rlI4'))
    order = {n: i for i, n in enumerate(NOTE)}
    def k(x):
        t = x['tom']; mm = re.match(r'([A-G][#b]?)(m?)', t)
        if not mm: return (99, 1, x['fonte'])
        r = FLAT.get(mm.group(1), mm.group(1))
        return (order.get(r, 98), 1 if mm.group(2) else 0, x['fonte'] != 'Incluído no app (arquivo de áudio)', x['nome'])
    return sorted(rows, key=k)

COLS_M = [('id','ID'),('nome','Nome padronizado (Música - Artista)'),('titulo','Música'),('artista','Artista'),('tom','Tom'),
          ('bpm','BPM'),('capo','Capotraste (casa)'),('pad','Pad'),('cifra','Tem letra e cifra no app'),
          ('arquivo','Arquivo de origem'),('slide','Slide do PowerPoint'),('repo','Arquivo no repositório'),('obs','Repetidas / observações')]
COLS_P = [('tom','Tom do pad'),('nome','Nome'),('link','Ouvir (link)'),('fonte','Fonte'),('offline','Funciona offline'),('arquivo','Arquivo no repositório')]

def main():
    os.makedirs(OUT, exist_ok=True)
    S, P = songs(), pads()
    for name, cols, rows in (('musicas.csv', COLS_M, S), ('pads.csv', COLS_P, P)):
        with open(os.path.join(OUT, name), 'w', newline='', encoding='utf-8') as fh:
            w = csv.writer(fh); w.writerow([c[1] for c in cols])
            for r in rows: w.writerow([r[c[0]] for c in cols])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.formatting.rule import FormulaRule
    wb = Workbook()
    for idx, (tab, cols, rows, widths) in enumerate((
        ('Músicas', COLS_M, S, [6, 46, 34, 26, 7, 7, 10, 8, 12, 30, 10, 44, 70]),
        ('Pads', COLS_P, P, [10, 34, 62, 32, 10, 22]))):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = tab
        ws.append([c[1] for c in cols])
        for r in rows:
            ws.append([r[c[0]] if r[c[0]] != '' else None for c in cols])
        for i, wdt in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = wdt
        hdr = PatternFill('solid', fgColor='111111')
        for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr; c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 30; ws.freeze_panes = 'C2' if tab == 'Músicas' else 'B2'
        last = f'{chr(64 + len(cols))}{len(rows) + 1}'
        t = Table(displayName='Musicas' if idx == 0 else 'Pads', ref=f'A1:{last}')
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True); ws.add_table(t)
        if tab == 'Pads':
            for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
                for c in row:
                    if c.value: c.hyperlink = c.value; c.font = Font(color='1155CC', underline='single')
        else:
            for row in ws.iter_rows(min_row=2, min_col=13, max_col=13):
                for c in row: c.alignment = Alignment(wrap_text=False)
            red = PatternFill('solid', fgColor='FCE4E4'); yel = PatternFill('solid', fgColor='FFF4CC')
            ws.conditional_formatting.add(f'A2:{last}', FormulaRule(formula=['ISNUMBER(SEARCH("REPETIDA",$M2))'], fill=red))
            ws.conditional_formatting.add(f'I2:I{len(rows)+1}', FormulaRule(formula=['$I2="Não"'], fill=yel))
    wb.save(os.path.join(OUT, 'Worship Church - Controle.xlsx'))
    rep = sum(1 for s in S if 'REPETIDA' in s['obs']); com = sum(1 for s in S if s['cifra'] == 'Sim')
    print(f'{len(S)} músicas ({com} com cifra, {rep} marcadas como REPETIDA), {len(P)} pads -> controle/')

if __name__ == '__main__':
    main()
